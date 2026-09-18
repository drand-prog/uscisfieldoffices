#!/usr/bin/env python3
"""
Ingest USCIS N-400 field-office performance reports into the dashboard's
JSON dataset (data/n400_quarterly.json).

Usage:
    python3 scripts/ingest_quarterly.py data/raw/n400_performance_data_fy2026_q4_v1.xlsx [more files...]
    python3 scripts/ingest_quarterly.py --all          # re-ingest everything in data/raw/

Accepts .xlsx, .csv, or .pdf source files, each covering one fiscal quarter
(filename must contain fyYYYY_qN or fyYYYYqtrN). USCIS has changed this
report's column layout many times since FY2015 -- see scripts/parse_common.py
and scripts/parse_pdf.py for how each era is handled. Re-running this script
for a quarter already in the dataset overwrites that quarter in place.
"""
import json
import re
import sys
from pathlib import Path

import openpyxl

sys.path.insert(0, str(Path(__file__).resolve().parent))
from parse_common import CODE_RE, iter_grid_offices, values_to_blocks  # noqa: E402
from parse_pdf import iter_pdf_offices  # noqa: E402

REPO_ROOT = Path(__file__).resolve().parent.parent
RAW_DIR = REPO_ROOT / "data" / "raw"
OUTPUT_PATH = REPO_ROOT / "data" / "n400_quarterly.json"

QUARTER_MONTHS = {
    1: ((10, 1), (12, 31)),
    2: ((1, 1), (3, 31)),
    3: ((4, 1), (6, 30)),
    4: ((7, 1), (9, 30)),
}

FILENAME_RE = re.compile(r"fy(\d{4}).{0,6}?q(?:tr)?([1-4])", re.IGNORECASE)


def quarter_dates(fy, q):
    (sm, sd), (em, ed) = QUARTER_MONTHS[q]
    year = fy - 1 if q == 1 else fy
    return f"{year:04d}-{sm:02d}-{sd:02d}", f"{year:04d}-{em:02d}-{ed:02d}"


def parse_fy_quarter_from_filename(path: Path):
    m = FILENAME_RE.search(path.name)
    if not m:
        raise ValueError(f"Can't find fyYYYY_qN in filename {path.name!r}")
    return int(m.group(1)), int(m.group(2))


def normalize_name(s):
    return re.sub(r"\s+", " ", (s or "").strip()).lower()


def load_grid(path: Path):
    if path.suffix.lower() == ".xlsx":
        wb = openpyxl.load_workbook(path, data_only=True)
        ws = wb[wb.sheetnames[0]]
        return [[c.value for c in row] for row in ws.iter_rows()]
    elif path.suffix.lower() == ".csv":
        import csv

        for encoding in ("utf-8-sig", "cp1252"):
            try:
                with open(path, encoding=encoding, newline="") as f:
                    return list(csv.reader(f))
            except UnicodeDecodeError:
                continue
        raise UnicodeDecodeError("csv", b"", 0, 1, f"could not decode {path.name}")
    raise ValueError(f"load_grid: unsupported extension {path.suffix}")


def build_name_code_lookup(raw_dir: Path):
    """Scan every spreadsheet-era file for (name, state)->code pairs, so quarters
    that lack a code column (FY2015 Q1) can still be keyed by stable codes.
    Keyed by (name, state) first since a handful of office names repeat across
    states (e.g. "Portland" in both Maine and Oregon); falls back to name-only
    when that's unambiguous."""
    by_name_state = {}
    by_name = {}
    ambiguous_names = set()
    for path in sorted(raw_dir.glob("*")):
        if path.suffix.lower() not in (".xlsx", ".csv"):
            continue
        try:
            grid = load_grid(path)
            for rec in iter_grid_offices(grid):
                if rec["code"] and rec["code"] != "TOTAL" and rec["name"]:
                    norm_name = normalize_name(rec["name"])
                    norm_state = normalize_name(rec["state"] or "")
                    by_name_state.setdefault((norm_name, norm_state), rec["code"])
                    prior = by_name.setdefault(norm_name, rec["code"])
                    if prior != rec["code"]:
                        ambiguous_names.add(norm_name)
        except Exception as exc:  # noqa: BLE001
            print(f"  (lookup pass skipped {path.name}: {exc})", file=sys.stderr)
    for name in ambiguous_names:
        del by_name[name]
    return by_name_state, by_name


def ingest_file(path: Path, by_name_state: dict, by_name: dict):
    fy, q = parse_fy_quarter_from_filename(path)
    quarter_key = f"FY{fy}Q{q}"
    start_date, end_date = quarter_dates(fy, q)
    ext = path.suffix.lower()

    offices = {}
    unresolved = []

    if ext == ".pdf":
        source_format = "pdf"
        row_iter = iter_pdf_offices(path)
    else:
        source_format = ext.lstrip(".")
        row_iter = iter_grid_offices(load_grid(path))

    for rec in row_iter:
        code = rec["code"]
        if code is None:
            key = (normalize_name(rec["name"]), normalize_name(rec["state"] or ""))
            code = by_name_state.get(key) or by_name.get(normalize_name(rec["name"]))
            if code is None:
                unresolved.append(rec["name"])
                continue
        blocks, suppressed = values_to_blocks(rec["raw_values"])
        offices[code] = {
            "code": code,
            "name": "All Field Offices (National)" if code == "TOTAL" else rec["name"].strip(),
            "state": rec["state"],
            "nat": blocks["nat"],
            "mil": blocks["mil"],
            "total": blocks["total"],
            "suppressed": suppressed,
        }

    if unresolved:
        print(f"  WARNING {path.name}: could not resolve code for: {unresolved}", file=sys.stderr)

    return quarter_key, {
        "fy": fy,
        "quarter": q,
        "label": f"FY{fy} Q{q}",
        "periodStart": start_date,
        "periodEnd": end_date,
        "sourceFile": path.name,
        "sourceFormat": source_format,
        "needsVerification": source_format == "pdf",
        "offices": offices,
    }


def main(argv):
    if not argv:
        print(__doc__)
        return 1

    if argv == ["--all"]:
        paths = sorted(RAW_DIR.glob("n400_performance_data_*"))
    else:
        paths = [Path(a) for a in argv]

    print("Building field-office name->code lookup from spreadsheet-era files...")
    by_name_state, by_name = build_name_code_lookup(RAW_DIR)
    print(f"  {len(by_name_state)} (name, state) pairs, {len(by_name)} unambiguous names")

    if OUTPUT_PATH.exists():
        dataset = json.loads(OUTPUT_PATH.read_text())
    else:
        dataset = {"quarters": {}}

    for path in paths:
        if not path.exists():
            print(f"skip (not found): {path}", file=sys.stderr)
            continue
        quarter_key, quarter_data = ingest_file(path, by_name_state, by_name)
        dataset["quarters"][quarter_key] = quarter_data
        n_offices = len(quarter_data["offices"]) - 1
        flag = " [PDF-derived]" if quarter_data["needsVerification"] else ""
        print(f"ingested {path.name} -> {quarter_key} ({n_offices} field offices){flag}")

    ordered = dict(
        sorted(dataset["quarters"].items(), key=lambda kv: kv[1]["fy"] * 10 + kv[1]["quarter"])
    )
    dataset["quarters"] = ordered
    dataset["generatedNote"] = (
        "Generated by scripts/ingest_quarterly.py. Do not hand-edit; re-run the "
        "script against the source workbook(s) instead."
    )

    OUTPUT_PATH.parent.mkdir(parents=True, exist_ok=True)
    OUTPUT_PATH.write_text(json.dumps(dataset, indent=2, sort_keys=False) + "\n")
    print(f"wrote {OUTPUT_PATH} ({len(dataset['quarters'])} quarters)")
    return 0


if __name__ == "__main__":
    raise SystemExit(main(sys.argv[1:]))
