"""
Shared parsing helpers for USCIS N-400 quarterly performance reports.

USCIS has changed this report's layout many times since FY2015: column
offsets shift, the "Field Office Code" column is sometimes unlabeled or
missing entirely (FY2015), and older files are PDFs instead of spreadsheets.
Rather than hard-coding row/column numbers per era, the grid parser below
*locates* the TOTAL row and the field-office-code column empirically from
each file's own content, then reads every row the same way regardless of
which era it came from.
"""
import re
from pathlib import Path

QUARTER_MONTHS = {
    1: ((10, 1), (12, 31)),
    2: ((1, 1), (3, 31)),
    3: ((4, 1), (6, 30)),
    4: ((7, 1), (9, 30)),
}

FILENAME_RE = re.compile(r"fy(\d{4}).{0,6}?q(?:tr)?([1-4])", re.IGNORECASE)


def quarter_dates(fy, q):
    """The USCIS federal fiscal quarter's real calendar start/end dates."""
    (sm, sd), (em, ed) = QUARTER_MONTHS[q]
    year = fy - 1 if q == 1 else fy
    return f"{year:04d}-{sm:02d}-{sd:02d}", f"{year:04d}-{em:02d}-{ed:02d}"


def parse_fy_quarter_from_filename(path: Path):
    m = FILENAME_RE.search(path.name)
    if not m:
        raise ValueError(f"Can't find fyYYYY_qN in filename {path.name!r}")
    return int(m.group(1)), int(m.group(2))


def normalize_name(s):
    return re.sub(r"\s+", " ", (s or "").strip()).lower()


def load_grid(path: Path):
    """Load a spreadsheet (.xlsx or .csv) into a grid of cell values."""
    if path.suffix.lower() == ".xlsx":
        import openpyxl

        wb = openpyxl.load_workbook(path, data_only=True)
        ws = wb[wb.sheetnames[0]]
        return [[c.value for c in row] for row in ws.iter_rows()]
    elif path.suffix.lower() == ".csv":
        import csv

        for encoding in ("utf-8-sig", "cp1252"):
            try:
                with open(path, encoding=encoding, newline="") as f:
                    return list(csv.reader(f))
            except UnicodeDecodeError:
                continue
        raise UnicodeDecodeError("csv", b"", 0, 1, f"could not decode {path.name}")
    raise ValueError(f"load_grid: unsupported extension {path.suffix}")


FOOTER_PREFIXES = (
    "table key",
    "references",
    "notes",
    "note:",
    "source",
    "d  disclosure",
    "d disclosure",
    "d data withheld",
    "- represents zero",
    "international field office",
    "international",
)

STRUCTURAL_LABELS = {
    "field office by state",
    "field office by territory",
    "uscis field office location",
    "uscis field office or service center location",
    "applications by category and case status",
    "usciss field office or service center",
}

CODE_RE = re.compile(r"^[A-Z]{2,5}$")

# US states/territories the report groups offices under. A row exactly
# matching one of these is a section header, even if (as happens in at least
# one source file, FY2017 Q4) it was corrupted to also carry leftover
# total-like figures -- those figures are never trusted as office data.
KNOWN_STATE_NAMES = {
    "alabama", "alaska", "arizona", "arkansas", "california", "colorado",
    "connecticut", "delaware", "district of columbia", "florida", "georgia",
    "guam", "hawaii", "idaho", "illinois", "indiana", "iowa", "kansas",
    "kentucky", "louisiana", "maine", "maryland", "massachusetts", "michigan",
    "minnesota", "mississippi", "missouri", "montana", "nebraska", "nevada",
    "new hampshire", "new jersey", "new mexico", "new york", "north carolina",
    "north dakota", "northern mariana islands", "ohio", "oklahoma", "oregon",
    "pennsylvania", "puerto rico", "rhode island", "south carolina",
    "south dakota", "tennessee", "texas", "utah", "vermont", "virginia",
    "virgin islands", "u.s. virgin islands", "washington", "west virginia",
    "wisconsin", "wyoming",
}


def is_known_state_name(raw_name):
    return normalize_label(raw_name) in KNOWN_STATE_NAMES


def normalize_label(s):
    s = (s or "").strip().lower()
    s = re.sub(r"\d+$", "", s).strip()
    return s


def looks_like_footer(raw_name):
    if raw_name is None:
        return True
    stripped = raw_name.strip()
    if not stripped:
        return False
    if re.match(r"^\d+[\s).]", stripped):
        return True
    norm = normalize_label(stripped)
    return any(norm.startswith(p) for p in FOOTER_PREFIXES)


def is_structural_label(raw_name):
    return normalize_label(raw_name) in STRUCTURAL_LABELS


def parse_number(raw):
    """Return (value_or_None, suppressed_bool) for one data cell."""
    if raw is None:
        return None, False
    if isinstance(raw, (int, float)):
        return raw, False
    s = str(raw).strip()
    if s == "":
        return None, False
    if s.upper() == "D":
        return None, True
    if s.upper() == "N/A":
        return None, False
    # "Represents zero" -- USCIS source files use several different dash
    # characters for this across eras/exports (ASCII hyphen, em dash, and at
    # least one PDF using U+2010 HYPHEN instead of ASCII '-').
    if s in ("-", "‐", "‑", "‒", "–", "—", "―"):
        return 0, False
    s2 = s.replace(",", "")
    return float(s2) if "." in s2 else int(s2), False


def is_number_ish(raw):
    if raw is None:
        return False
    if isinstance(raw, (int, float)):
        return True
    s = str(raw).strip()
    if s == "":
        return False
    try:
        parse_number(s)
        return True
    except ValueError:
        return False


def cell_text(v):
    return "" if v is None else str(v).strip()


def find_total_row(grid):
    """Return (row_idx, values_start_col) for the row whose col0 label is TOTAL/Grand Total."""
    for r, row in enumerate(grid):
        label = cell_text(row[0] if row else None)
        if label.lower() in ("total", "grand total"):
            values_start_col = None
            for c in range(1, len(row)):
                if is_number_ish(row[c]):
                    values_start_col = c
                    break
            if values_start_col is not None:
                return r, values_start_col
    raise ValueError("Could not locate TOTAL row in sheet")


def detect_code_col(grid, total_row_idx, values_start_col, lookahead=60):
    if values_start_col - 1 < 0:
        return None
    candidate = values_start_col - 1
    seen = 0
    matches = 0
    for r in range(total_row_idx + 1, min(len(grid), total_row_idx + 1 + lookahead)):
        row = grid[r]
        if candidate >= len(row):
            continue
        text = cell_text(row[candidate])
        if not text:
            continue
        seen += 1
        if CODE_RE.match(text):
            matches += 1
    if seen == 0:
        return None
    return candidate if (matches / seen) > 0.5 else None


def row_name(row, boundary):
    for c in range(boundary - 1, -1, -1):
        if c < len(row):
            text = cell_text(row[c])
            if text:
                return text
    return ""


def iter_grid_offices(grid, num_values=12):
    """Yield dicts: {name, code (maybe None), state, raw_values (list of num_values)}
    plus a synthetic first entry with name='TOTAL' for the national total row.

    num_values is the number of data columns per row (4 per category block:
    received/approved/denied/pending) -- 12 for N-400's three blocks
    (Naturalization/Military/Total), 20 for I-485's five (Family/Employment/
    Humanitarian/Other/Total). Located the same way regardless: empirically,
    from wherever the TOTAL row's numbers actually start."""
    total_row_idx, values_start_col = find_total_row(grid)
    code_col = detect_code_col(grid, total_row_idx, values_start_col)
    boundary = code_col if code_col is not None else values_start_col

    total_row = grid[total_row_idx]
    total_values = [total_row[values_start_col + i] if values_start_col + i < len(total_row) else None for i in range(num_values)]
    yield {
        "name": "TOTAL",
        "code": "TOTAL",
        "state": None,
        "raw_values": total_values,
    }

    current_state = None
    for r in range(total_row_idx + 1, len(grid)):
        row = grid[r]
        if not row:
            continue
        name = row_name(row, boundary)
        if not name:
            continue
        if looks_like_footer(name):
            break
        if is_structural_label(name):
            continue
        if normalize_label(name) in ("total", "grand total"):
            # Some source files (at least one I-485 CSV era) repeat the
            # national total a second time near the end of the office list,
            # after the real one this function already yielded separately.
            continue

        raw_values = [row[values_start_col + i] if values_start_col + i < len(row) else None for i in range(num_values)]
        has_data = any(cell_text(v) not in ("", " ") for v in raw_values)

        code = None
        if code_col is not None and code_col < len(row):
            code_text = cell_text(row[code_col])
            if code_text:
                code = code_text

        if not has_data:
            current_state = name
            continue

        # A row named after a state/territory but with no code is a genuine
        # section header even if (as in at least one source file) it was
        # corrupted to also carry leftover total-like figures -- but a real
        # office literally named after its state (e.g. "Washington", "New
        # York") always carries its own code, so check that first. Only
        # applies to layouts that have a code column at all (FY2015 files
        # don't, and there every real office row is already distinguished
        # from a header purely by has_data).
        if code_col is not None and code is None and is_known_state_name(name):
            current_state = name
            continue

        yield {
            "name": name,
            "code": code,
            "state": current_state,
            "raw_values": raw_values,
        }


def values_to_blocks(raw_values):
    """Split 12 raw cell values into nat/mil/total blocks of {received,approved,denied,pending}."""
    return values_to_named_blocks(raw_values, ("nat", "mil", "total"))


def values_to_named_blocks(raw_values, bucket_names):
    """Split raw cell values into named blocks of {received,approved,denied,pending},
    4 values per bucket in bucket_names order (the general form of values_to_blocks,
    for report layouts with a different number/set of category blocks)."""
    keys = ("received", "approved", "denied", "pending")
    blocks = {}
    suppressed = {}
    for bi, bucket in enumerate(bucket_names):
        vals = {}
        sup = {}
        for ki, key in enumerate(keys):
            val, is_sup = parse_number(raw_values[bi * 4 + ki])
            vals[key] = val
            sup[key] = is_sup
        blocks[bucket] = vals
        suppressed[bucket] = sup
    return blocks, suppressed
